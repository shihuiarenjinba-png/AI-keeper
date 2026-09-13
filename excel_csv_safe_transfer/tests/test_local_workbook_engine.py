from pathlib import Path

from openpyxl import Workbook, load_workbook

from transfer.config import AppConfig
from transfer.csv_loader import load_csv
from transfer.local_workbook_engine import LocalWorkbookEngine, file_sha256


ROOT = Path(__file__).resolve().parents[1]


def test_in_place_transfer_preserves_formula_and_creates_exact_backup(tmp_path: Path):
    config = AppConfig.load(ROOT / "config_example.json")
    data_dir = tmp_path / "顧客データ (本社)-2026_09"
    data_dir.mkdir()
    csv_path = data_dir / "売上_東京-01.csv"
    excel_path = data_dir / "2026年度 売上台帳.xlsx"

    csv_path.write_text(
        "日付,伝票番号,数量,金額\n2026/09/01,00123,2,1500\n",
        encoding="utf-8-sig",
    )
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "入力"
    worksheet["B2"] = "2026年度"
    worksheet["E5"] = "=C5*D5"
    workbook.save(excel_path)
    workbook.close()

    original_hash = file_sha256(excel_path)
    result = LocalWorkbookEngine(config).process_in_place(
        load_csv(csv_path, config),
        excel_path,
        expected_sha256=original_hash,
    )

    assert Path(result.backup_path).is_file()
    assert file_sha256(result.backup_path) == original_hash
    assert result.sha256_before == original_hash
    assert result.sha256_after == file_sha256(excel_path)

    updated = load_workbook(excel_path, data_only=False)
    try:
        worksheet = updated["入力"]
        assert worksheet["B5"].value == "00123"
        assert worksheet["E5"].value == "=C5*D5"
    finally:
        updated.close()
